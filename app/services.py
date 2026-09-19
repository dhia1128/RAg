from __future__ import annotations

import httpx
from typing import List

from sqlalchemy.orm import Session

from app.config import OLLAMA_BASE_URL, OPENAI_MODEL
from app.models import Conversation, Document, DocumentChunk


class DocumentService:
    def __init__(self, db: Session):
        self.db = db

    def extract_text(self, file_content: bytes, file_name: str) -> str:
        if file_name.endswith(".txt"):
            return file_content.decode("utf-8", errors="ignore")

        if file_name.endswith(".pdf"):
            try:
                from pypdf import PdfReader
                import io

                reader = PdfReader(io.BytesIO(file_content))
                pages = []
                for page in reader.pages:
                    pages.append(page.extract_text() or "")
                return "\n".join(pages)
            except Exception as exc:
                raise ValueError(f"Failed to read PDF file: {exc}") from exc

        if file_name.endswith(".docx"):
            try:
                from docx import Document as DocxDocument
                import io

                document = DocxDocument(io.BytesIO(file_content))
                paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
                return "\n".join(paragraphs)
            except Exception as exc:
                raise ValueError(f"Failed to read DOCX file: {exc}") from exc

        raise ValueError("Unsupported file format. Please upload .txt, .pdf, or .docx.")

    def chunk_text(self, text: str, chunk_size: int = 800, overlap: int = 120) -> List[str]:
        if len(text) <= chunk_size:
            return [text.strip()] if text.strip() else []

        chunks: List[str] = []
        start = 0
        while start < len(text):
            end = start + chunk_size
            chunk = text[start:end]
            if chunk.strip():
                chunks.append(chunk.strip())
            if end >= len(text):
                break
            start += chunk_size - overlap
        return chunks

    def create_document(self, title: str, file_name: str, content_type: str, content: str) -> Document:
        document = Document(
            title=title,
            file_name=file_name,
            content_type=content_type,
            content=content,
        )
        self.db.add(document)
        self.db.flush()

        chunks = self.chunk_text(content)

        for index, chunk in enumerate(chunks):
            self.db.add(DocumentChunk(document_id=document.id, chunk_index=index, chunk_text=chunk))

        self.db.commit()
        self.db.refresh(document)
        return document

    def list_documents(self) -> List[Document]:
        return self.db.query(Document).order_by(Document.created_at.desc()).all()

    def ask_question(self, question: str, document_id: int | None = None) -> dict[str, str | int]:
        if document_id is None:
            documents = self.list_documents()
            if not documents:
                raise RuntimeError("No documents available yet.")
            document_id = documents[0].id

        document = self.db.query(Document).filter(Document.id == document_id).first()
        if document is None:
            raise RuntimeError("Document not found.")

        chunks = self.db.query(DocumentChunk).filter(DocumentChunk.document_id == document_id).all()
        if not chunks:
            raise RuntimeError("No chunks available for this document.")

        relevant_context = "\n\n".join(chunk.chunk_text for chunk in chunks[:5])

        answer = self._call_ollama(question=question, context=relevant_context)

        conversation = Conversation(
            document_id=document_id,
            user_question=question,
            ai_answer=answer,
        )
        self.db.add(conversation)
        self.db.commit()

        return {
            "document_id": document_id,
            "question": question,
            "answer": answer,
        }

    def _call_ollama(self, question: str, context: str) -> str:
        model_name = OPENAI_MODEL or "llama3.2"
        base_url = OLLAMA_BASE_URL.rstrip("/")

        try:
            response = httpx.post(
                f"{base_url}/api/generate",
                json={
                    "model": model_name,
                    "prompt": f"Answer the question using only the provided document context.\n\nContext:\n{context}\n\nQuestion:\n{question}",
                    "stream": False,
                },
                timeout=120.0,
            )
            response.raise_for_status()
            data = response.json()
            return data.get("response", "").strip() or "I could not generate a response from the local model."
        except Exception as exc:
            raise RuntimeError(
                "Could not reach local Ollama. Please install Ollama and run a model such as 'llama3.2'."
            ) from exc
